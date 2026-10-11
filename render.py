"""Write a report as .docx and .pdf on the hospital letterhead:

  [crest]   IBB SPECIALIST HOSPITAL, MINNA   [hospital logo]
            PMB 169 CHANGHAGA, NIGER STATE.
                    RADIOLOGY REPORT
  | Surname | Other Name (s) | Age | Sex | Hospital No | X-Ray/USS/CT No |
  report title, findings, conclusion (radiologist's house style: Times New Roman 12 pt)
  date                                              radiologist / title

Letterhead text and logos come from letterhead/letterhead.json and the image files next to it,
so they can be replaced without touching code. Unapproved reports carry a DRAFT banner.
"""
FILES_VERSION = "2026-10-11"   # must match streamlit_app.py

import io, json
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from xml.sax.saxutils import escape

DRAFT = "DRAFT – computer-generated, not yet reviewed and signed by the radiologist"
LH_DIR = Path(__file__).resolve().parent / 'letterhead'
TABLE_COLS = [("Surname", "surname", 2.25), ("Other Name (s)", "other_names", 2.3), ("Age", "age", 0.75),
              ("Sex", "sex_label", 0.7), ("Hospital No", "hospital_no", 1.6), ("X-Ray/USS/CT No", "study_no", 2.4)]
MARGIN_LR = 0.75   # inches; the letterhead and patient table use the full text width


def letterhead():
    try:
        cfg = json.loads((LH_DIR / 'letterhead.json').read_text(encoding='utf-8'))
    except Exception:
        return None
    for k in ('left_logo', 'right_logo'):
        p = LH_DIR / cfg.get(k, '') if cfg.get(k) else None
        cfg[k] = str(p) if p and p.exists() else None
    return cfg


def _patient_cells(r):
    return [str(r.get(key) or '') for _, key, _ in TABLE_COLS]


# =========================================================================== Word
def _cell_borders_none(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'nil'); b.append(e)
    tblPr.append(b)


def _fix_widths(table, widths):
    """Make Word/LibreOffice honour column widths (autofit off + grid widths + cell widths)."""
    table.autofit = False
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.gridCol_lst, widths):
        gc.w = w
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = w


def _run(p, text, font=None, size=None, bold=False, underline=False, color=None):
    run = p.add_run(text)
    run.bold, run.underline = bold, underline
    if font:
        run.font.name = font
        run._element.rPr.rFonts.set(qn('w:ascii'), font); run._element.rPr.rFonts.set(qn('w:hAnsi'), font)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def to_docx(r):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Inches(MARGIN_LR)
    s.top_margin, s.bottom_margin = Inches(0.5), Inches(0.7)
    st = d.styles['Normal']
    st.font.name, st.font.size = 'Times New Roman', Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(6)

    def para(text='', bold=False, underline=False, align=None, color=None, size=None, font=None, after=None):
        p = d.add_paragraph()
        if align is not None:
            p.alignment = align
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        _run(p, text, font=font, size=size, bold=bold, underline=underline, color=color)
        return p

    if not r.get('approved'):
        para(DRAFT, bold=True, color=(0xB0, 0x1E, 0x1E), size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)

    lh = letterhead()
    if lh:
        t = d.add_table(rows=1, cols=3)
        _cell_borders_none(t)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = [Inches(1.1), Inches(4.55), Inches(1.1)]
        _fix_widths(t, widths)
        for i in range(3):
            t.cell(0, i).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if lh.get('left_logo'):
            t.cell(0, 0).paragraphs[0].add_run().add_picture(lh['left_logo'], width=Inches(1.0))
        c = t.cell(0, 1)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
        _run(p, lh.get('hospital', ''), font='Arial Black', size=16, bold=True)
        p = c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
        _run(p, lh.get('address', ''), font='Arial', size=11)
        if lh.get('right_logo'):
            rp = t.cell(0, 2).paragraphs[0]; rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            rp.add_run().add_picture(lh['right_logo'], width=Inches(1.0))
        p = para(lh.get('heading', 'RADIOLOGY REPORT'), bold=True, font='Arial', size=14,
                 align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
        p.paragraph_format.space_before = Pt(8)

    # patient table
    pt = d.add_table(rows=2, cols=len(TABLE_COLS))
    pt.style = 'Table Grid'
    pt.alignment = WD_TABLE_ALIGNMENT.CENTER
    total = sum(w for *_, w in TABLE_COLS)
    usable = 8.27 - 2 * MARGIN_LR
    _fix_widths(pt, [Inches(usable * w / total) for *_, w in TABLE_COLS])
    for i, ((head, _, w), val) in enumerate(zip(TABLE_COLS, _patient_cells(r))):
        hp = pt.cell(0, i).paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(hp, head, font='Arial', size=10.5, bold=True)
        vp = pt.cell(1, i).paragraphs[0]
        _run(vp, val, font='Arial', size=10.5)
        for row in (0, 1):
            pt.cell(row, i).paragraphs[0].paragraph_format.space_after = Pt(3)
            pt.cell(row, i).paragraphs[0].paragraph_format.space_before = Pt(3)
    d.add_paragraph().paragraph_format.space_after = Pt(2)

    para(r['title'], bold=True, underline=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    if r.get('referrer'):
        p = d.add_paragraph(); _run(p, 'Referred by: ', bold=True); _run(p, r['referrer'])
    if r.get('indication'):
        p = d.add_paragraph(); _run(p, 'INDICATION: ', bold=True); _run(p, r['indication'], bold=True)
        if r.get('exam') == 'breast':
            para('FINDINGS:', bold=True)
    for t_ in r['paragraphs']:
        p = para(t_); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p = d.add_paragraph()
    _run(p, 'Conclusion: ', bold=True)
    concl = r['conclusion']
    if len(concl) == 1:
        _run(p, concl[0], bold=True)
    else:
        _run(p, 'Features are suggestive of:', bold=True)
        for i, c in enumerate(concl, 1):
            para(f"{i}. {c}", bold=True)
    for a in r.get('advice', []):
        para('NB: ' + a, bold=True)
    d.add_paragraph()
    sig = d.add_table(rows=2, cols=2)
    sig.cell(0, 0).text = r['date_str']
    sig.cell(0, 1).text = r['radiologist']
    sig.cell(1, 1).text = r['radiologist_title']
    for c in (sig.cell(0, 1), sig.cell(1, 1)):
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


# ============================================================================ PDF
def _img(path, width):
    from reportlab.lib.utils import ImageReader
    iw, ih = ImageReader(path).getSize()
    return Image(path, width=width, height=width * ih / iw)


def to_pdf(r):
    buf = io.BytesIO()
    W = A4[0] - 2 * MARGIN_LR * inch
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=MARGIN_LR * inch, rightMargin=MARGIN_LR * inch,
                            topMargin=0.5 * inch, bottomMargin=0.7 * inch, title=r['title'], author=r['radiologist'])
    base = ParagraphStyle('b', fontName='Times-Roman', fontSize=12, leading=15, spaceAfter=7, alignment=TA_JUSTIFY)
    bold = ParagraphStyle('bb', parent=base, fontName='Times-Bold', alignment=0)
    title = ParagraphStyle('t', parent=bold, alignment=TA_CENTER, spaceAfter=10)
    draft = ParagraphStyle('d', parent=bold, fontSize=9, leading=11, alignment=TA_CENTER, textColor=colors.HexColor('#B01E1E'), spaceAfter=2)
    hosp = ParagraphStyle('h', fontName='Helvetica-Bold', fontSize=17, leading=21, alignment=TA_CENTER)
    addr = ParagraphStyle('a', fontName='Helvetica', fontSize=11, leading=13, alignment=TA_CENTER)
    head = ParagraphStyle('hd', fontName='Helvetica-Bold', fontSize=14, leading=17, alignment=TA_CENTER, spaceBefore=8, spaceAfter=10)
    th = ParagraphStyle('th', fontName='Helvetica-Bold', fontSize=10, leading=12, alignment=TA_CENTER)
    td = ParagraphStyle('td', fontName='Helvetica', fontSize=10, leading=12)
    E = escape
    story = []
    if not r.get('approved'):
        story.append(Paragraph(E(DRAFT), draft))
    lh = letterhead()
    if lh:
        lw = 0.95 * inch
        left = _img(lh['left_logo'], lw) if lh.get('left_logo') else ''
        right = _img(lh['right_logo'], lw) if lh.get('right_logo') else ''
        mid = [Paragraph(E(lh.get('hospital', '')), hosp), Paragraph(E(lh.get('address', '')), addr)]
        t = Table([[left, mid, right]], colWidths=[1.05 * inch, W - 2.1 * inch, 1.05 * inch])
        t.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('ALIGN', (0, 0), (0, 0), 'LEFT'),
                               ('ALIGN', (2, 0), (2, 0), 'RIGHT'), ('LEFTPADDING', (0, 0), (-1, -1), 0),
                               ('RIGHTPADDING', (0, 0), (-1, -1), 0)]))
        story += [t, Paragraph(E(lh.get('heading', 'RADIOLOGY REPORT')), head)]
    total = sum(w for *_, w in TABLE_COLS)
    pt = Table([[Paragraph(E(h), th) for h, *_ in TABLE_COLS], [Paragraph(E(v), td) for v in _patient_cells(r)]],
               colWidths=[W * w / total for *_, w in TABLE_COLS])
    pt.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.8, colors.black), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                            ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
    story += [pt, Spacer(1, 14), Paragraph(f"<u>{E(r['title'])}</u>", title)]
    if r.get('referrer'):
        story.append(Paragraph(f"<b>Referred by:</b> {E(r['referrer'])}", ParagraphStyle('rf', parent=base, alignment=0)))
    if r.get('indication'):
        story.append(Paragraph(f"INDICATION: {E(r['indication'])}", bold))
        if r.get('exam') == 'breast':
            story.append(Paragraph("FINDINGS:", bold))
    for t_ in r['paragraphs']:
        story.append(Paragraph(E(t_), base))
    concl = r['conclusion']
    if len(concl) == 1:
        story.append(Paragraph(f"Conclusion: {E(concl[0])}", bold))
    else:
        story.append(Paragraph("Conclusion: Features are suggestive of:", bold))
        for i, c in enumerate(concl, 1):
            story.append(Paragraph(f"{i}. {E(c)}", bold))
    for a in r.get('advice', []):
        story.append(Paragraph(f"NB: {E(a)}", bold))
    story.append(Spacer(1, 24))
    sig = Table([[r['date_str'], r['radiologist']], ['', r['radiologist_title']]], colWidths=[W / 2, W / 2])
    sig.setStyle(TableStyle([('FONT', (0, 0), (-1, -1), 'Times-Roman', 12), ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
                             ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0)]))
    story.append(sig)
    doc.build(story)
    return buf.getvalue()
